import os
import csv
import sqlite3
import shutil

def export_all_to_csv(output_dir="dataset_csv"):
    base_dir = os.path.dirname(os.path.abspath(__file__))
    db_path = os.path.join(base_dir, "seafood_compliance.db")
    project_root = os.path.abspath(os.path.join(base_dir, ".."))
    target_dir = os.path.join(project_root, output_dir)
    os.makedirs(target_dir, exist_ok=True)

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
    tables = [row[0] for row in cursor.fetchall() if not row[0].startswith("sqlite")]

    exported_files = {}

    print(f"Exporting database tables from {db_path} to {target_dir}...")

    for table in tables:
        csv_file_path = os.path.join(target_dir, f"{table}.csv")
        cursor.execute(f"SELECT * FROM {table}")
        rows = cursor.fetchall()
        if rows:
            headers = rows[0].keys()
            with open(csv_file_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                for row in rows:
                    writer.writerow(list(row))
            exported_files[table] = {"path": csv_file_path, "count": len(rows)}
            print(f"  [+] {table}.csv: {len(rows)} records")
        else:
            print(f"  [-] {table}: 0 records (skipped)")

    # Generate unified Master Denormalized CSV dataset
    master_csv_path = os.path.join(target_dir, "seafood_compliance_master_dataset.csv")
    root_master_csv_path = os.path.join(project_root, "master_dataset.csv")
    
    master_query = """
    SELECT 
        sl.log_id,
        sl.timestamp AS log_timestamp,
        sl.sensor_id,
        s.sensor_type,
        s.battery_status AS sensor_battery_pct,
        s.signal_status AS sensor_signal,
        s.calibration_status,
        s.accuracy_rating AS sensor_accuracy,
        sl.batch_id,
        pb.product_type,
        pb.quantity_kg,
        pb.origin AS batch_origin,
        pb.destination AS batch_destination,
        pb.required_temp_min,
        pb.required_temp_max,
        sl.shipment_id,
        shp.vehicle_id,
        shp.driver_id,
        shp.port_airport,
        shp.shipment_status,
        sl.temperature AS recorded_temp,
        sl.humidity AS recorded_humidity,
        sl.location AS gps_location,
        sl.is_missing,
        sl.is_noisy,
        sl.is_imputed,
        sl.original_temp,
        sl.imputed_temp,
        sl.anomaly_score,
        sl.sensor_status,
        pb.compliance_status AS batch_compliance,
        shp.risk_score AS shipment_risk_score
    FROM sensor_logs sl
    LEFT JOIN shipments shp ON sl.shipment_id = shp.shipment_id
    LEFT JOIN product_batches pb ON sl.batch_id = pb.batch_id
    LEFT JOIN sensors s ON sl.sensor_id = s.sensor_id
    ORDER BY sl.log_id ASC
    """
    cursor.execute(master_query)
    master_rows = cursor.fetchall()
    if master_rows:
        headers = master_rows[0].keys()
        for path in [master_csv_path, root_master_csv_path]:
            with open(path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                for row in master_rows:
                    writer.writerow(list(row))
        exported_files["master_dataset_csv"] = {"path": root_master_csv_path, "count": len(master_rows)}
        print(f"  [+] master_dataset.csv created with {len(master_rows)} records")

    # Generate multi-sheet master_dataset.xlsx
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        header_font = Font(name='Calibri', bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
        header_align = Alignment(horizontal='center', vertical='center')

        def add_sheet(title, query):
            ws = wb.create_sheet(title=title)
            cursor.execute(query)
            q_rows = cursor.fetchall()
            if q_rows:
                q_headers = q_rows[0].keys()
                ws.append(list(q_headers))
                for cell in ws[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_align
                for r in q_rows:
                    ws.append(list(r))
                for col_idx in range(1, len(q_headers) + 1):
                    ws.column_dimensions[get_column_letter(col_idx)].width = 18

        # Sheet 1: Master Combined (active)
        ws_master = wb.active
        if ws_master is None:
            ws_master = wb.create_sheet("Master Compliance Dataset")
        else:
            ws_master.title = "Master Compliance Dataset"
        cursor.execute(master_query + " LIMIT 5000")
        m_rows = cursor.fetchall()
        if m_rows:
            m_headers = m_rows[0].keys()
            ws_master.append(list(m_headers))
            for cell in ws_master[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
            for r in m_rows:
                ws_master.append(list(r))
            for col_idx in range(1, len(m_headers) + 1):
                ws_master.column_dimensions[get_column_letter(col_idx)].width = 18

        add_sheet("Shipments", "SELECT * FROM shipments")
        add_sheet("Product Batches", "SELECT * FROM product_batches")
        add_sheet("Sensor Calibrations", "SELECT * FROM sensor_calibrations")
        add_sheet("Custody Handovers", "SELECT * FROM handover_records")
        add_sheet("Route Events", "SELECT * FROM route_events")
        add_sheet("Worker Workloads", "SELECT * FROM worker_logs")
        add_sheet("Compliance Events", "SELECT * FROM compliance_events")
        add_sheet("Sensor Logs (Sample)", "SELECT * FROM sensor_logs LIMIT 5000")

        root_xlsx_path = os.path.join(project_root, "master_dataset.xlsx")
        wb.save(root_xlsx_path)
        print(f"  [+] master_dataset.xlsx created at {root_xlsx_path}")
        exported_files["master_dataset_xlsx"] = {"path": root_xlsx_path}
    except Exception as e:
        print(f"  [-] Excel export error: {e}")

    conn.close()
    print("Dataset Artifacts Generation Complete!")
    return exported_files

if __name__ == "__main__":
    export_all_to_csv()
